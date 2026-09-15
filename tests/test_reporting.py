from pathlib import Path

from openpyxl import load_workbook

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "output"

REQUIRED_SHEETS = [
    "Research Summary", "Company Financials", "REIT Metrics", "DCF Valuation",
    "Comparable Companies", "Portfolio Analytics", "Scenario Analysis",
    "Data Quality", "Assumptions and Data Dictionary",
]

EXPECTED_HEADERS = {
    "Company Financials": {"ticker", "fiscal_year", "revenue_mm", "ebitda_mm", "net_income_mm"},
    "REIT Metrics": {"ticker", "property_value_mm", "noi_mm", "ffo_mm", "affo_mm", "cap_rate"},
    "DCF Valuation": {"ticker", "dcf_value_per_share", "enterprise_value", "equity_value"},
    "Comparable Companies": {"ticker", "peer_group", "implied_value_per_share"},
    "Data Quality": {"table", "issue_type", "ticker", "field", "detail"},
}


def test_excel_pack_exists(full_pipeline_run):
    assert (OUTPUT_DIR / "investment_research_pack.xlsx").exists()


def test_excel_pack_has_required_sheets(full_pipeline_run):
    wb = load_workbook(OUTPUT_DIR / "investment_research_pack.xlsx")
    for sheet in REQUIRED_SHEETS:
        assert sheet in wb.sheetnames


def test_excel_pack_sheets_have_expected_columns(full_pipeline_run):
    wb = load_workbook(OUTPUT_DIR / "investment_research_pack.xlsx")
    for sheet_name, expected_cols in EXPECTED_HEADERS.items():
        ws = wb[sheet_name]
        found_headers = set()
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            row_values = {str(c.value) for c in row if c.value is not None}
            if expected_cols & row_values:
                found_headers |= row_values
        missing = expected_cols - found_headers
        assert not missing, f"{sheet_name} is missing expected columns: {missing}"


def test_excel_pack_has_charts(full_pipeline_run):
    wb = load_workbook(OUTPUT_DIR / "investment_research_pack.xlsx")
    total_charts = sum(len(wb[s]._charts) for s in wb.sheetnames)
    assert total_charts >= 4


def test_research_note_written(full_pipeline_run):
    notes = list(OUTPUT_DIR.glob("research_note_*.md"))
    assert len(notes) >= 1
    text = notes[0].read_text()
    assert "Simulated analysis" in text
    assert "Bull / Base / Bear" in text
